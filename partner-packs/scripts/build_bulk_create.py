#!/usr/bin/env python3
"""Write the Canva Bulk Create workbook — one sheet per pack page, one row per
sector, headers matching the text slots in the Canva template.

  python scripts/build_bulk_create.py
  python scripts/build_bulk_create.py --out data/SlugABug_BulkCreate.xlsx

Sheet 1 keeps the exact headers of the incoming
data/SlugABug_BulkCreate_AllSectors_v4.csv so an existing Bulk Create
connection on the cover page keeps working.
"""

import argparse
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import engine

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DEFAULT_OUT = os.path.join(ROOT, "out", "SlugABug_BulkCreate_PartnerPacks.xlsx")

HEAD_FILL = PatternFill("solid", fgColor="1E1E1E")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)


def _n(items, label, count):
    return ["%s%d" % (label, i + 1) for i in range(count)]


def sheet_cover(c):
    return [("Industry", c["industry"]), ("Tag", c["tag"]),
            ("Title", "PARTNER PACK"), ("Subhead", c["cover_subhead"]),
            ("Body", c["cover_body"])] + [
        kv for i, (t, b) in enumerate(c["cover_tiles"])
        for kv in (("Tile%d_Title" % (i + 1), t), ("Tile%d_Body" % (i + 1), b))
    ] + [("Footer", c["footer_strap"])]


def sheet_service_overview(c):
    rows = [("Industry", c["industry"]),
            ("Heading_Red", c["so_title_red"]), ("Heading_Ink", c["so_title_ink"]),
            ("Body", c["so_body"])]
    for i, (t, b, _) in enumerate(c["so_tiles"]):
        rows += [("Tile%d_Title" % (i + 1), t.replace("&amp;", "&")),
                 ("Tile%d_Body" % (i + 1), b)]
    for i, a in enumerate(c["areas11"]):
        rows.append(("Area%d" % (i + 1), a))
    for i, (t, b, _) in enumerate(c["program"]):
        rows += [("Step%d_Title" % (i + 1), t), ("Step%d_Body" % (i + 1), b)]
    rows += [("Footer_Lead", c["so_footer_lead"]),
             ("Footer_Tail", c["so_footer_tail"])]
    return rows


def sheet_sample_report(c):
    rows = [("Industry", c["industry"]), ("Badge", "EXAMPLE ONLY"),
            ("Property", c["sr_property"]), ("Date", c["sr_date"]),
            ("ServiceType", c["sr_service_type"]),
            ("Technician", c["sr_technician"]), ("SiteContact", c["sr_contact"]),
            ("JobNo", c["sr_job_no"])]
    for label, key in (("Area", "sr_areas"), ("Finding", "sr_findings"),
                       ("Action", "sr_actions"), ("Rec", "sr_recs")):
        for i, v in enumerate(c[key]):
            rows.append(("%s%d" % (label, i + 1), v))
    rows += [("NextVisit", c["sr_next_visit"]), ("Footer", c["footer_strap"])]
    return rows


def sheet_case_study(c):
    rows = [("Industry", c["industry"]), ("Subhead", c["case_subhead"]),
            ("Intro", c["case_intro"])]
    for label, key in (("Challenge", "challenges"), ("Solution", "solutions"),
                       ("Result", "results"), ("WhyMatters", "why_matters")):
        for i, v in enumerate(c[key]):
            rows.append(("%s%d" % (label, i + 1), v))
    rows += [("Quote", c["quote"]), ("Attribution", c["attribution"]),
             ("CTA", "%s %s" % (c["cta_lead"], c["cta_tail"]))]
    return rows


def sheet_quote_scope(c):
    rows = [("Industry", c["industry"]), ("Subhead", c["quote_subhead"]),
            ("SiteLabel", c["site_label"])]
    for label, key in (("Scope", "scope6"), ("Access", "access6"),
                       ("Area", "areas9")):
        for i, v in enumerate(c[key]):
            rows.append(("%s%d" % (label, i + 1), v))
    rows += [("CTA_Lead", "Ready to price your site?"),
             ("CTA_Tail", "%s %s" % (c["cta_lead"], c["cta_tail"]))]
    return rows


def sheet_prevention(c):
    rows = [("Industry", c["industry"]), ("Title_Red", c["pg_title_red"]),
            ("Title_Ink", "PEST PREVENTION GUIDE"), ("Subhead", c["pg_subhead"])]
    for i, t in enumerate(c["tips"]):
        rows.append(("Tip%d" % (i + 1), t))
    for i, (a, b) in enumerate(c["urgent"]):
        rows.append(("Urgent%d" % (i + 1), "%s\n%s" % (a, b)))
    rows += [("NeedHelp", c["need_help"]), ("Footer", c["footer_strap"])]
    return rows


def sheet_capability(c):
    rows = [("Industry", c["industry"]), ("Overview", c["cap_overview"])]
    for i, v in enumerate(c["core"]):
        rows.append(("Core%d" % (i + 1), v))
    for i, v in enumerate(c["snapshot"]):
        rows.append(("Snapshot%d" % (i + 1), v))
    for i, v in enumerate(c["why"]):
        rows.append(("Why%d" % (i + 1), v))
    for i, (t, b, _) in enumerate(c["cap_tiles"]):
        rows += [("Tile%d_Title" % (i + 1), t.replace("<br>", " ")),
                 ("Tile%d_Body" % (i + 1), b)]
    rows.append(("CTA", "%s %s" % (c["cta_lead"], c["cta_tail"])))
    return rows


SHEETS = [
    ("1_Cover", sheet_cover),
    ("2_ServiceOverview", sheet_service_overview),
    ("3_SampleReport", sheet_sample_report),
    ("4_CaseStudy", sheet_case_study),
    ("5_QuoteScope", sheet_quote_scope),
    ("6_PreventionGuide", sheet_prevention),
    ("7_Capability", sheet_capability),
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=DEFAULT_OUT)
    args = ap.parse_args()

    contexts = engine.build_all()
    wb = Workbook()
    wb.remove(wb.active)

    for name, fn in SHEETS:
        ws = wb.create_sheet(name)
        rows = [fn(c) for c in contexts]
        headers = [k for k, _ in rows[0]]
        ws.append(headers)
        for row in rows:
            ws.append([v for _, v in row])
        for i, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=i)
            cell.fill = HEAD_FILL
            cell.font = HEAD_FONT
            cell.alignment = Alignment(vertical="center")
            width = max(len(h) + 2, min(46, max(
                (len(str(r[i - 1][1])) for r in rows), default=12) + 2))
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "B2"

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    wb.save(args.out)
    print("wrote %s  (%d sheets x %d sectors)"
          % (os.path.relpath(args.out, ROOT), len(SHEETS), len(contexts)))


if __name__ == "__main__":
    main()
